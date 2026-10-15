"""CV Markdown -> Word (.docx, python-docx), pour les ATS qui lisent mieux le Word que le PDF.

Mêmes règles que ``cv_pdf`` : une colonne, styles Word natifs (``Heading 1-3``, ``List Bullet``) que les ATS
reconnaissent, coordonnées dans le corps du document, Arial, aucune image ni tableau.
"""

from __future__ import annotations

import io
from collections.abc import Iterable

from . import cv_format

_FONT = "Arial"
_MUTED = (90, 90, 90)
_HEADING_SIZES = {1: 18, 2: 12, 3: 11}


def render_docx(
    markdown: str, title: str = "CV", *, author: str = "", keywords: Iterable[str] = (), lang: str = "fr"
) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(16)

    def set_font(style, size: float) -> None:
        style.font.name, style.font.size = _FONT, Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)  # sinon la police du thème l'emporte sur Arial

    set_font(doc.styles["Normal"], 10)
    for level, size in _HEADING_SIZES.items():
        set_font(doc.styles[f"Heading {level}"], size)

    props = doc.core_properties
    props.title, props.subject, props.author = title[:120], title[:120], author
    props.keywords, props.language = ", ".join(keywords), lang

    def paragraph(text: str, style: str | None = None, *, size: float | None = None, muted: bool = False):
        par = doc.add_paragraph(style=style)
        par.paragraph_format.space_after = Pt(2)
        for part, span_style, _url in cv_format.spans(text):
            run = par.add_run(part)
            run.bold = "B" in span_style or None
            run.italic = "I" in span_style or None
            if size:
                run.font.size = Pt(size)
            if muted:
                run.font.color.rgb = RGBColor(*_MUTED)
        return par

    in_header = True
    for block in cv_format.blocks(markdown):
        if block.kind == "blank":
            continue
        if block.kind == "heading":
            level = min(block.level, 3)
            in_header = in_header and level == 1
            text = block.text.upper() if level == 2 else block.text
            doc.add_heading(cv_format.plain(text), level=level)
        elif block.kind == "bullet":
            paragraph(block.text, "List Bullet")
        elif in_header and block.text.startswith("**"):
            paragraph(block.text, size=12)
        elif in_header or (block.text.startswith("*") and not block.text.startswith("**")):
            paragraph(block.text, size=9.5, muted=True)
        else:
            paragraph(block.text)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
